"""
export_service.py — Design kit export generators.
Produces .md, .pdf, .docx, and .zip exports using Jinja2 templates.
Handles Unicode-safe PDF generation and resilient ZIP bundling.
"""
import io
import json
import logging
import re
import zipfile
from pathlib import Path

from jinja2 import Environment, FileSystemLoader

logger = logging.getLogger(__name__)

TEMPLATE_DIR = Path(__file__).parent.parent.parent / "templates" / "exports"
env = Environment(loader=FileSystemLoader(str(TEMPLATE_DIR)), autoescape=False)


def safe_filename_slug(name: str, fallback: str = "export", max_len: int = 30) -> str:
    """Build an ASCII-only, filename-safe slug from an arbitrary name.

    Strips everything except ``[a-z0-9-]``, collapses runs of dashes, and
    trims to *max_len* characters.  Returns *fallback* if the result is empty.
    """
    slug = re.sub(r'[^a-z0-9]+', '-', (name or "").lower()).strip('-')[:max_len]
    return slug or fallback


def _sanitize_latin1(text: str) -> str:
    """Sanitize text for Latin-1 encoding (used by fpdf2 built-in fonts).

    Replaces common Unicode characters that Claude generates with Latin-1 equivalents.
    Falls back to encoding with replace for any remaining non-Latin-1 characters.
    """
    if not text:
        return ""
    replacements = {
        '\u2014': '--',   # em-dash
        '\u2013': '-',    # en-dash
        '\u2018': "'",    # left single quote
        '\u2019': "'",    # right single quote
        '\u201c': '"',    # left double quote
        '\u201d': '"',    # right double quote
        '\u2026': '...',  # ellipsis
        '\u2022': '*',    # bullet
        '\u2023': '>',    # triangular bullet
        '\u00a0': ' ',    # non-breaking space
        '\u200b': '',     # zero-width space
        '\u2192': '->',   # right arrow
        '\u2190': '<-',   # left arrow
        '\u2194': '<->',  # left-right arrow
        '\u2713': '[x]',  # check mark
        '\u2717': '[ ]',  # cross mark
        '\u00b7': '*',    # middle dot
        '\u2028': '\n',   # line separator
        '\u2029': '\n',   # paragraph separator
    }
    for char, replacement in replacements.items():
        text = text.replace(char, replacement)
    # Final fallback: encode to latin-1 with replace to catch any remaining
    return text.encode('latin-1', errors='replace').decode('latin-1')


def _build_context(sheet, blocks, pipeline_nodes) -> dict:
    """Build template context from project artifacts."""
    return {
        "problem": sheet.problem or "Not specified",
        "audience": sheet.audience or "Not specified",
        "mvp": sheet.mvp or "Not specified",
        "features": sheet.features or [],
        "tone": sheet.tone or "Not specified",
        "platform": sheet.platform or "Not specified",
        "tech_constraints": sheet.tech_constraints or "None",
        "success_metric": sheet.success_metric or "Not specified",
        "confidence_score": sheet.confidence_score or 0,
        "blocks": [
            {
                "name": b.name,
                "description": b.description or "",
                "category": b.category,
                "priority": b.priority,
                "effort": b.effort,
                "is_mvp": b.is_mvp,
            }
            for b in blocks
        ],
        "mvp_blocks": [b for b in blocks if b.is_mvp],
        "v2_blocks": [b for b in blocks if not b.is_mvp],
        "pipeline": [
            {"layer": n.layer, "tool": n.selected_tool}
            for n in pipeline_nodes
        ],
    }


def build_export_context_from_artifact(artifact_ctx: dict) -> dict:
    """Build template context from the unified artifact context dict.

    Produces the same dict shape as ``_build_context()`` so every generator
    works unchanged.  For v2 projects the discovery summary is included
    as an extra key for template branching.
    """
    blocks_orm = artifact_ctx.get("blocks", [])
    pipeline_orm = artifact_ctx.get("pipeline", [])

    block_dicts = [
        {
            "name": b.name,
            "description": b.description or "",
            "category": b.category,
            "priority": b.priority,
            "effort": b.effort,
            "is_mvp": b.is_mvp,
        }
        for b in blocks_orm
    ]

    return {
        "project_name": artifact_ctx.get("project_name", ""),
        "flow_version": artifact_ctx.get("flow_version", "v1"),
        "problem": artifact_ctx.get("problem") or "Not specified",
        "audience": artifact_ctx.get("audience") or "Not specified",
        "mvp": artifact_ctx.get("mvp") or "Not specified",
        "features": artifact_ctx.get("features") or [],
        "tone": artifact_ctx.get("tone") or "Not specified",
        "platform": artifact_ctx.get("platform") or "Not specified",
        "tech_constraints": artifact_ctx.get("tech_constraints") or "None",
        "success_metric": artifact_ctx.get("success_metric") or "Not specified",
        "confidence_score": artifact_ctx.get("confidence_score") or 0,
        "discovery_summary": artifact_ctx.get("discovery_summary") or "",
        "modules": artifact_ctx.get("modules", []),
        "blocks": block_dicts,
        "mvp_blocks": [b for b in block_dicts if b.get("is_mvp")],
        "v2_blocks": [b for b in block_dicts if not b.get("is_mvp")],
        "pipeline": [
            {"layer": n.layer, "tool": n.selected_tool}
            for n in pipeline_orm
        ],
    }


async def generate_markdown(sheet=None, blocks=None, pipeline_nodes=None, *, context=None) -> str:
    """Generate a Markdown export of the project design kit."""
    if context is None:
        context = _build_context(sheet, blocks, pipeline_nodes)

    try:
        template = env.get_template("design_kit.md.j2")
        return template.render(**context)
    except Exception:
        # Fallback: generate inline
        is_v2 = context.get("flow_version") == "v2"
        lines = ["# Design Kit", ""]

        if is_v2 and context.get("discovery_summary"):
            lines.extend([
                "## Discovery Insights",
                context["discovery_summary"],
                "",
            ])
            if context.get("audience") and context["audience"] != "Not specified":
                lines.extend(["## Target Audience", context["audience"], ""])
        else:
            lines.extend([
                "## Problem",
                f"{context['problem']}",
                "",
                "## Target Audience",
                f"{context['audience']}",
                "",
                "## MVP Scope",
                f"{context['mvp']}",
                "",
            ])

        lines.append("## Features")
        for b in context["blocks"]:
            lines.append(f"- **{b['name']}** ({b['priority'].upper()}, {b['effort']}) -- {b['description']}")

        lines.extend(["", "## Tech Stack"])
        for p in context["pipeline"]:
            lines.append(f"- {p['layer']}: {p['tool']}")

        if not is_v2:
            lines.extend([
                "",
                "## Constraints",
                f"{context['tech_constraints']}",
                "",
                "## Success Metric",
                f"{context['success_metric']}",
            ])

        confidence = context.get("confidence_score", 0)
        footer = "*Generated by Ide/AI*" if is_v2 else f"*Generated by Ide/AI | Confidence: {confidence}%*"
        lines.extend(["", "---", footer])
        return "\n".join(lines)


async def generate_text(sheet=None, blocks=None, pipeline_nodes=None, *, context=None) -> str:
    """Generate a plain text export."""
    if context is None:
        context = _build_context(sheet, blocks, pipeline_nodes)
    md = await generate_markdown(context=context)
    # Strip markdown formatting for plain text
    text = md.replace("# ", "").replace("## ", "").replace("**", "").replace("*", "")
    text = text.replace("---", "=" * 40)
    return text


async def generate_pdf(sheet=None, blocks=None, pipeline_nodes=None, *, context=None) -> bytes:
    """Generate a PDF export of the project design kit using fpdf2.

    Uses Latin-1 sanitization for built-in Helvetica font compatibility.
    """
    from fpdf import FPDF

    if context is None:
        context = _build_context(sheet, blocks, pipeline_nodes)
    s = _sanitize_latin1  # shorthand
    is_v2 = context.get("flow_version") == "v2"

    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()

    # Title
    pdf.set_font("Helvetica", "B", 24)
    pdf.cell(0, 15, "Design Kit", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(5)

    if is_v2 and context.get("discovery_summary"):
        # v2: discovery summary as main content
        pdf.set_font("Helvetica", "B", 14)
        pdf.cell(0, 10, "Discovery Insights", new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", "", 11)
        for line in context["discovery_summary"].split("\n"):
            line = line.strip()
            if not line:
                pdf.ln(3)
            elif line.startswith("## "):
                pdf.set_font("Helvetica", "B", 12)
                pdf.cell(0, 8, s(line[3:]), new_x="LMARGIN", new_y="NEXT")
                pdf.set_font("Helvetica", "", 11)
            elif line.startswith("- "):
                pdf.multi_cell(0, 6, s(f"  {line}"))
            else:
                pdf.multi_cell(0, 6, s(line))
        pdf.ln(5)
    else:
        # v1: sheet fields
        pdf.set_font("Helvetica", "B", 14)
        pdf.cell(0, 10, "Problem", new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", "", 11)
        pdf.multi_cell(0, 6, s(context["problem"]))
        pdf.ln(5)

        pdf.set_font("Helvetica", "B", 14)
        pdf.cell(0, 10, "Target Audience", new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", "", 11)
        pdf.multi_cell(0, 6, s(context["audience"]))
        pdf.ln(5)

        pdf.set_font("Helvetica", "B", 14)
        pdf.cell(0, 10, "MVP Scope", new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", "", 11)
        pdf.multi_cell(0, 6, s(context["mvp"]))
        pdf.ln(5)

    # Features
    pdf.set_font("Helvetica", "B", 14)
    pdf.cell(0, 10, "Features", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 11)
    for b in context["blocks"]:
        text = f"  * {b['name']} ({b['priority'].upper()}, {b['effort']}) - {b['description']}"
        pdf.multi_cell(0, 6, s(text))
    if not context["blocks"]:
        pdf.multi_cell(0, 6, "No features defined yet.")
    pdf.ln(5)

    # Tech Stack
    pdf.set_font("Helvetica", "B", 14)
    pdf.cell(0, 10, "Tech Stack", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 11)
    for p in context["pipeline"]:
        pdf.cell(0, 6, s(f"  * {p['layer']}: {p['tool']}"), new_x="LMARGIN", new_y="NEXT")
    if not context["pipeline"]:
        pdf.cell(0, 6, "No tech stack defined yet.", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(5)

    if not is_v2:
        # Constraints
        pdf.set_font("Helvetica", "B", 14)
        pdf.cell(0, 10, "Constraints", new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", "", 11)
        pdf.multi_cell(0, 6, s(context["tech_constraints"]))
        pdf.ln(5)

        # Success Metric
        pdf.set_font("Helvetica", "B", 14)
        pdf.cell(0, 10, "Success Metric", new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", "", 11)
        pdf.multi_cell(0, 6, s(context["success_metric"]))
        pdf.ln(5)

    # Footer
    pdf.set_font("Helvetica", "I", 9)
    footer = "Generated by Ide/AI" if is_v2 else f"Generated by Ide/AI | Confidence: {context['confidence_score']}%"
    pdf.cell(0, 10, footer)

    return pdf.output()


async def generate_docx(sheet=None, blocks=None, pipeline_nodes=None, *, context=None) -> bytes:
    """Generate a DOCX export of the project design kit using python-docx."""
    from docx import Document

    if context is None:
        context = _build_context(sheet, blocks, pipeline_nodes)
    is_v2 = context.get("flow_version") == "v2"

    doc = Document()

    doc.add_heading("Design Kit", level=0)

    if is_v2 and context.get("discovery_summary"):
        doc.add_heading("Discovery Insights", level=1)
        for line in context["discovery_summary"].split("\n"):
            line = line.strip()
            if not line:
                continue
            if line.startswith("## "):
                doc.add_heading(line[3:], level=2)
            elif line.startswith("- "):
                doc.add_paragraph(line[2:], style="List Bullet")
            else:
                doc.add_paragraph(line)
    else:
        doc.add_heading("Problem", level=1)
        doc.add_paragraph(context["problem"])

        doc.add_heading("Target Audience", level=1)
        doc.add_paragraph(context["audience"])

        doc.add_heading("MVP Scope", level=1)
        doc.add_paragraph(context["mvp"])

    doc.add_heading("Features", level=1)
    if context["blocks"]:
        for b in context["blocks"]:
            p = doc.add_paragraph(style="List Bullet")
            run = p.add_run(f"{b['name']} ")
            run.bold = True
            p.add_run(f"({b['priority'].upper()}, {b['effort']}) -- {b['description']}")
    else:
        doc.add_paragraph("No features defined yet.")

    doc.add_heading("Tech Stack", level=1)
    if context["pipeline"]:
        for p_item in context["pipeline"]:
            doc.add_paragraph(f"{p_item['layer']}: {p_item['tool']}", style="List Bullet")
    else:
        doc.add_paragraph("No tech stack defined yet.")

    if not is_v2:
        doc.add_heading("Constraints", level=1)
        doc.add_paragraph(context["tech_constraints"])

        doc.add_heading("Success Metric", level=1)
        doc.add_paragraph(context["success_metric"])

    # Footer
    footer = doc.add_paragraph()
    footer_text = "Generated by Ide/AI" if is_v2 else f"Generated by Ide/AI | Confidence: {context['confidence_score']}%"
    footer.add_run(footer_text).italic = True

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


async def generate_zip(sheet=None, blocks=None, pipeline_nodes=None, *, context=None) -> bytes:
    """Generate a ZIP export containing all available format exports.

    Resilient: if PDF or DOCX generation fails, the ZIP still includes MD and TXT.
    Failed formats get an error note file instead.
    """
    if context is None:
        context = _build_context(sheet, blocks, pipeline_nodes)
    md_content = await generate_markdown(context=context)
    txt_content = await generate_text(context=context)

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("design-kit.md", md_content)
        zf.writestr("design-kit.txt", txt_content)

        # PDF — may fail on unusual Unicode
        try:
            pdf_content = await generate_pdf(context=context)
            zf.writestr("design-kit.pdf", pdf_content)
        except Exception as e:
            logger.warning("PDF generation failed in ZIP bundle: %s", e)
            zf.writestr("design-kit-pdf-error.txt", f"PDF generation failed: {e}\n\nPlease export PDF separately after reviewing your content for special characters.")

        # DOCX — may fail if python-docx encounters issues
        try:
            docx_content = await generate_docx(context=context)
            zf.writestr("design-kit.docx", docx_content)
        except Exception as e:
            logger.warning("DOCX generation failed in ZIP bundle: %s", e)
            zf.writestr("design-kit-docx-error.txt", f"DOCX generation failed: {e}\n\nPlease export DOCX separately.")

    return buf.getvalue()
